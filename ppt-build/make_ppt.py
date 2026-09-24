"""Build SCGIP presentation PPTX (16:9) — 1M1B × IBM SkillsBuild guideline structure."""
import os, shutil
from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

PASTE = "C:/Users/lithe/AppData/Local/Temp/freebuff-desktop-pastes"
HERE = os.path.dirname(os.path.abspath(__file__))
IMG1 = os.path.join(PASTE, "paste-1789274572970-1536.png")   # citizen portal (wide)
IMG2 = os.path.join(PASTE, "paste-1789274720864-1536.png")   # ward map (portrait)
IMG4 = os.path.join(PASTE, "paste-1789281243089-1536.png")   # notification center (wide)

# ---- palette ----
GREEN_D = RGBColor(0x0B, 0x3D, 0x2E)   # deep forest
GREEN   = RGBColor(0x14, 0x53, 0x2D)
EMER    = RGBColor(0x10, 0xB9, 0x81)   # emerald accent
AMBER   = RGBColor(0xF5, 0x9E, 0x0B)
LIGHT   = RGBColor(0xF4, 0xF7, 0xF5)
INK     = RGBColor(0x1F, 0x29, 0x37)
MUT     = RGBColor(0x5B, 0x6B, 0x63)
WHITE   = RGBColor(0xFF, 0xFF, 0xFF)
TEAL    = RGBColor(0x0E, 0x74, 0x90)

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]
FONT = "Calibri"

def new_slide(bgcolor=LIGHT):
    s = prs.slides.add_slide(BLANK)
    r = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    r.fill.solid(); r.fill.fore_color.rgb = bgcolor
    r.line.fill.background(); r.shadow.inherit = False
    return s

def rect(s, x, y, w, h, fill, rounded=False, line=None):
    shp = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
                             Inches(x), Inches(y), Inches(w), Inches(h))
    shp.fill.solid(); shp.fill.fore_color.rgb = fill
    if line: shp.line.color.rgb = line; shp.line.width = Pt(1)
    else: shp.line.fill.background()
    shp.shadow.inherit = False
    return shp

def tbox(s, x, y, w, h, anchor=MSO_ANCHOR.TOP):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    return tf

def para(tf, text, size=14, bold=False, color=INK, align=PP_ALIGN.LEFT,
         before=0, after=6, first=False, italic=False, font=FONT):
    p = tf.paragraphs[0] if (first and not tf.paragraphs[0].runs) else tf.add_paragraph()
    p.alignment = align; p.space_before = Pt(before); p.space_after = Pt(after)
    r = p.add_run(); r.text = text
    f = r.font; f.name = font; f.size = Pt(size); f.bold = bold; f.italic = italic; f.color.rgb = color
    return p

def header(s, title, kicker=None):
    rect(s, 0.55, 0.62, 0.09, 0.62, EMER)
    tf = tbox(s, 0.82, 0.50, 11.9, 1.0)
    if kicker:
        para(tf, kicker.upper(), size=11, bold=True, color=EMER, first=True, after=2)
        para(tf, title, size=27, bold=True, color=GREEN_D, after=0)
    else:
        para(tf, title, size=29, bold=True, color=GREEN_D, first=True, after=0)

def footer(s, n):
    tf = tbox(s, 9.3, 7.08, 3.6, 0.35)
    para(tf, f"SCGIP  ·  1M1B × IBM SkillsBuild  ·  {n:02d}", size=9, color=MUT,
         align=PP_ALIGN.RIGHT, first=True, after=0)

def pic(s, path, x, y, w=None, h=None, frame=True):
    im = Image.open(path); ar = im.size[0] / im.size[1]
    if w and not h: h = w / ar
    if h and not w: w = h * ar
    if frame:
        rect(s, x - 0.04, y - 0.04, w + 0.08, h + 0.08, WHITE)
    p = s.shapes.add_picture(path, Inches(x), Inches(y), Inches(w), Inches(h))
    return w, h

def caption(s, x, y, w, text):
    tf = tbox(s, x, y, w, 0.3)
    para(tf, text, size=10.5, italic=True, color=MUT, align=PP_ALIGN.CENTER, first=True, after=0)

def chip(s, x, y, w, h, text, fill=EMER, tcolor=WHITE, size=13, bold=True):
    c = rect(s, x, y, w, h, fill, rounded=True)
    tf = c.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    para(tf, text, size=size, bold=bold, color=tcolor, align=PP_ALIGN.CENTER, first=True, after=0)
    return c

def bullets(tf, items, size=14.5, gap=8, color=INK, bold_first=False):
    for i, it in enumerate(items):
        if bold_first and ":" in it:
            head, rest = it.split(":", 1)
            p = tf.paragraphs[0] if (i == 0 and not tf.paragraphs[0].runs) else tf.add_paragraph()
            p.space_after = Pt(gap)
            r1 = p.add_run(); r1.text = "▸ " + head + ":"
            r1.font.name = FONT; r1.font.size = Pt(size); r1.font.bold = True; r1.font.color.rgb = GREEN
            r2 = p.add_run(); r2.text = rest
            r2.font.name = FONT; r2.font.size = Pt(size); r2.font.color.rgb = color
        else:
            para(tf, "▸  " + it, size=size, color=color, first=(i == 0), after=gap)

# ============ SLIDE 1 — TITLE ============
s = new_slide(GREEN_D)
rect(s, 0, 0, 13.333, 0.18, EMER)
rect(s, 0, 7.32, 13.333, 0.18, EMER)
tf = tbox(s, 0.9, 1.15, 11.5, 1.4)
para(tf, "SCGIP", size=64, bold=True, color=WHITE, first=True, after=0)
para(tf, "Sanitation Co-Governance Intelligence Platform", size=24, bold=True, color=EMER, after=0)
tf = tbox(s, 0.9, 3.15, 11.5, 0.7)
para(tf, "An AI-Driven Real-Time Civic Complaint System for Urban Sanitation", size=18, color=WHITE, first=True)
chip(s, 0.9, 4.05, 4.6, 0.52, "🌍  UN SDG 6 — Clean Water & Sanitation", fill=EMER, tcolor=GREEN_D, size=14)
chip(s, 5.7, 4.05, 3.4, 0.52, "Real-Time · Co-Governed", fill=GREEN, tcolor=WHITE, size=13)
chip(s, 9.3, 4.05, 2.6, 0.52, "Born at B2G Hackathon", fill=GREEN, tcolor=WHITE, size=13)
tf = tbox(s, 0.9, 5.35, 11.5, 1.6)
para(tf, "Lithesh S  ·  Sri Sairam Institute of Technology, Chennai", size=17, bold=True, color=WHITE, first=True, after=4)
para(tf, "1M1B AI for Sustainability Virtual Internship — in collaboration with IBM SkillsBuild & AICTE",
     size=13, color=RGBColor(0xB9, 0xD6, 0xC9), after=2)
para(tf, "September 2026", size=12, color=RGBColor(0x8F, 0xB8, 0xA4), after=0)

# ============ SLIDE 2 — PROBLEM ============
s = new_slide(); header(s, "Reports Go Into a Void", "Problem Statement"); footer(s, 2)
tf = tbox(s, 0.82, 1.75, 6.6, 3.6)
bullets(tf, [
    "Citizens spot illegal dumping, blocked drains, overflowing garbage — but have no reliable way to report with evidence.",
    "Complaints by phone or in person vanish: no ticket number, no status, no proof anything happened.",
    "The Greater Chennai Corporation cannot prioritise hundreds of scattered reports or verify which are genuine.",
    "Result: recurring sanitation hazards, polluted drains and public spaces — and eroding citizen trust.",
], size=15, gap=12)
c = rect(s, 0.82, 5.55, 11.7, 1.25, GREEN_D, rounded=True)
tf = c.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
para(tf, "How might we use AI to score, route and track citizen sanitation reports transparently —",
     size=15.5, bold=True, color=WHITE, align=PP_ALIGN.CENTER, first=True, after=2)
para(tf, "so that citizens and the municipal corporation can verify and resolve issues together?",
     size=15.5, bold=True, color=EMER, align=PP_ALIGN.CENTER, after=0)

# ============ SLIDE 3 — SDG ALIGNMENT ============
s = new_slide(); header(s, "SDG Alignment", "Sustainability Focus"); footer(s, 3)
cards = [
    ("SDG 6", "Clean Water & Sanitation", "PRIMARY", EMER,
     "Direct action on blocked drains, illegal dumping and public-space sanitation — protecting water systems and community health."),
    ("SDG 11", "Sustainable Cities & Communities", "SECONDARY", TEAL,
     "Accountable, citizen-inclusive urban governance with measurable response deadlines."),
    ("SDG 16", "Peace, Justice & Strong Institutions", "SECONDARY", AMBER,
     "Transparency and auditability that build trust in civic institutions."),
]
x = 0.82
for tag, name, kind, accent, desc in cards:
    rect(s, x, 1.85, 3.75, 4.6, WHITE)
    rect(s, x, 1.85, 3.75, 0.16, accent)
    tf = tbox(s, x + 0.25, 2.15, 3.25, 4.1)
    para(tf, tag, size=30, bold=True, color=accent, first=True, after=2)
    para(tf, name, size=15.5, bold=True, color=GREEN_D, after=4)
    para(tf, kind, size=10.5, bold=True, color=MUT, after=10)
    para(tf, desc, size=13, color=INK, after=0)
    x += 3.98
tf = tbox(s, 0.82, 6.7, 11.7, 0.5)
para(tf, "One primary focus — SDG 6 — with direct secondary contributions to SDG 11 and SDG 16.",
     size=13, italic=True, color=MUT, align=PP_ALIGN.CENTER, first=True, after=0)

# ============ SLIDE 4 — DESIGN THINKING / ORIGIN ============
s = new_slide(); header(s, "From B2G Hackathon Prototype to Working Platform", "Design Thinking Journey"); footer(s, 4)
stages = [
    ("EMPATHIZE", "Citizens never see outcomes; officers receive unstructured, unprioritised reports."),
    ("DEFINE", "The real gap is verified closure — nobody confirms the fix, so nobody reports twice."),
    ("IDEATE", "Co-governance: place the citizen inside the workflow as the final verifier."),
    ("PROTOTYPE", "B2G Hackathon prototype → several weeks of iteration into a full real-time platform."),
    ("TEST & REFINE", "End-to-end live pipeline tested across citizen, officer and admin roles."),
]
y = 1.8
for i, (tag, desc) in enumerate(stages):
    accent = EMER if i % 2 == 0 else TEAL
    chip(s, 0.82, y, 2.35, 0.62, f"{i+1}.  {tag}", fill=accent, size=12.5)
    tf = tbox(s, 3.45, y + 0.02, 9.1, 0.85, anchor=MSO_ANCHOR.MIDDLE)
    para(tf, desc, size=14, color=INK, first=True, after=0)
    y += 0.98

# ============ SLIDE 5 — SOLUTION OVERVIEW ============
s = new_slide(); header(s, "SCGIP — Co-Governance by Design", "Solution Overview"); footer(s, 5)
steps = [
    ("1", "Citizen submits complaint — photo + real-time GPS location"),
    ("2", "SERI engine scores severity 0–100 (explainable)"),
    ("3", "Auto-routes to the right department with an SLA deadline (24h–15 days)"),
    ("4", "GCC ticket auto-generated and pushed to Corporation officers"),
    ("5", "Officer verifies → assigns → resolves, with photo evidence"),
    ("6", "Citizen confirms the fix — loop closed by both sides"),
]
y = 1.8
for n, txt in steps:
    chip(s, 0.82, y, 0.5, 0.5, n, fill=GREEN_D, size=15)
    tf = tbox(s, 1.55, y - 0.02, 5.9, 0.9, anchor=MSO_ANCHOR.MIDDLE)
    para(tf, txt, size=13.5, color=INK, first=True, after=0)
    y += 0.83
w, h = pic(s, IMG1, 7.75, 1.95, w=5.0)
caption(s, 7.75, 1.95 + h + 0.12, w, "Citizen portal — reporting dashboard with live data")
tf = tbox(s, 0.82, 6.75, 6.6, 0.6)
para(tf, "Three role-based portals: Citizen · Corporation Officer · Administrator",
     size=13, bold=True, color=GREEN, first=True, after=0)

# ============ SLIDE 6 — AI ELEMENTS ============
s = new_slide(); header(s, "AI Elements Used", "Where AI Powers the Platform"); footer(s, 6)
tf = tbox(s, 0.82, 1.7, 7.0, 5.4)
bullets(tf, [
    "SERI Explainable Severity Engine: complaint type × ward risk index × evidence quality — every score shows \"why this severity\".",
    "Agentic AI Assistant (Python FastAPI): tool-calling agent answering natural-language queries over live complaint data — grounded, never fabricated.",
    "NLP Sentiment Analysis: urgency signals detected from complaint descriptions.",
    "Geospatial Intelligence: GPS pins with reverse geocoding + ward-risk heat mapping on a live map.",
    "Real-time backbone: Server-Sent Events notifications, SLA auto-escalation, tamper-evident audit log.",
], size=14.5, gap=13)
c = rect(s, 8.15, 1.85, 4.4, 3.1, WHITE)
rect(s, 8.15, 1.85, 4.4, 0.14, AMBER)
tf = tbox(s, 8.4, 2.12, 3.9, 2.7)
para(tf, "IBM BOB", size=20, bold=True, color=GREEN_D, first=True, after=4)
para(tf, "Used during the ideation stage to refine the problem framing and design the co-governance workflow.",
     size=13, color=INK, after=8)
para(tf, "Prompt engineering guided both the ideation and the assistant's query logic.",
     size=12.5, italic=True, color=MUT, after=0)
c = rect(s, 8.15, 5.15, 4.4, 1.6, GREEN_D, rounded=True)
tf = c.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
para(tf, "\"An agent is only as smart as the context you feed it — grounding beats guessing.\"",
     size=13.5, italic=True, bold=True, color=EMER, align=PP_ALIGN.CENTER, first=True, after=0)

# ============ SLIDE 7 — PLATFORM WALKTHROUGH ============
s = new_slide(); header(s, "Platform Walkthrough", "Working Prototype — Screenshots"); footer(s, 7)
w, h = pic(s, IMG4, 0.82, 1.62, w=7.1)
caption(s, 0.82, 1.62 + h + 0.08, w, "Real-time notification center — every role stays in sync instantly")
y2 = 1.62 + h + 0.45
w1, h1 = pic(s, IMG1, 0.82, y2, w=4.35)
caption(s, 0.82, y2 + h1 + 0.08, w1, "Citizen portal — report with photo + GPS")
w2, h2 = pic(s, IMG2, 5.45, y2, h=h1)
caption(s, 5.3, y2 + h2 + 0.08, w2 + 0.3, "Live ward view — complaint pins")
tf = tbox(s, 8.3, 1.75, 4.5, 5.0)
bullets(tf, [
    "Live ward map with GPS pins, photos and reverse-geocoded addresses.",
    "Officer inbox with SLA countdown badges and priority queue.",
    "Complaint detail: SERI score, \"why this severity\", full timeline.",
    "Citizen confirms the fix — co-governance loop closed.",
], size=13.5, gap=12)

# ============ SLIDE 8 — TECHNOLOGY ARCHITECTURE ============
s = new_slide(); header(s, "Technology Architecture", "Built With"); footer(s, 8)
layers = [
    ("FRONTEND", "Next.js · TypeScript · Tailwind CSS · Leaflet + OpenStreetMap (live ward map)", EMER),
    ("REALTIME", "Server-Sent Events (SSE) — instant notifications across all roles · SLA auto-escalation", TEAL),
    ("BACKEND & DATA", "Next.js API routes · Prisma ORM · SQLite · GCC webhook ticket gateway", GREEN),
    ("AI SERVICES", "SERI engine + NLP sentiment (TypeScript) · Agentic AI assistant (Python FastAPI)", AMBER),
    ("SECURITY", "bcrypt password hashing · cookie-session auth · role-based access · tamper-evident audit log", RGBColor(0xB4, 0x50, 0x36)),
]
y = 1.85
for tag, txt, accent in layers:
    rect(s, 0.82, y, 11.7, 0.86, WHITE)
    rect(s, 0.82, y, 0.14, 0.86, accent)
    tf = tbox(s, 1.15, y + 0.06, 3.1, 0.75, anchor=MSO_ANCHOR.MIDDLE)
    para(tf, tag, size=13, bold=True, color=accent, first=True, after=0)
    tf = tbox(s, 4.3, y + 0.06, 8.0, 0.75, anchor=MSO_ANCHOR.MIDDLE)
    para(tf, txt, size=13, color=INK, first=True, after=0)
    y += 1.0

# ============ SLIDE 9 — TARGET USERS ============
s = new_slide(); header(s, "Target Users", "Who It Serves"); footer(s, 9)
users = [
    ("🏙️", "Citizens of Chennai", "Report issues with photo + GPS evidence, track progress live, and confirm the fix — closing the loop.", EMER),
    ("👮", "GCC Ward Officers & Sanitation Inspectors", "Receive an AI-prioritised queue with SLA deadlines; verify, assign and resolve with photo evidence.", TEAL),
    ("📊", "City Administrators", "Live ward map, SLA-compliance analytics and ward-level sanitation intelligence for data-driven decisions.", AMBER),
]
x = 0.82
for emoji, name, desc, accent in users:
    rect(s, x, 1.9, 3.75, 4.35, WHITE)
    rect(s, x, 1.9, 3.75, 0.14, accent)
    tf = tbox(s, x + 0.25, 2.2, 3.25, 3.9)
    para(tf, emoji, size=34, first=True, after=6)
    para(tf, name, size=16, bold=True, color=GREEN_D, after=8)
    para(tf, desc, size=13, color=INK, after=0)
    x += 3.98
tf = tbox(s, 0.82, 6.55, 11.7, 0.5)
para(tf, "Designed as a replicable blueprint — adoptable by any municipal body.",
     size=13, italic=True, color=MUT, align=PP_ALIGN.CENTER, first=True, after=0)

# ============ SLIDE 10 — EXPECTED IMPACT ============
s = new_slide(); header(s, "Expected Impact", "What Changes"); footer(s, 10)
tf = tbox(s, 0.82, 1.8, 11.6, 5.0)
bullets(tf, [
    "Closed complaint loop: every report is tracked from submission to citizen-confirmed resolution — no more complaints into a void.",
    "Faster, fairer response: severity-based prioritisation with SLA deadlines and auto-escalation replaces first-come-first-served handling.",
    "Accountability by design: a tamper-evident audit trail makes every action verifiable by citizens and administrators.",
    "Data-driven prevention: ward-level risk analytics reveal recurring hotspots — proactive action instead of reactive cleanup.",
    "Citizen trust: live visibility of the process itself builds trust in civic institutions.",
    "Scalable blueprint: the co-governance model extends to any ward, any city, any municipal body.",
], size=15.5, gap=15)
c = rect(s, 0.82, 6.35, 11.7, 0.75, GREEN_D, rounded=True)
tf = c.text_frame; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
para(tf, "From \"reported\" to \"resolved and verified\" — accountability as a system feature, not a promise.",
     size=14, bold=True, color=EMER, align=PP_ALIGN.CENTER, first=True, after=0)

# ============ SLIDE 11 — RESPONSIBLE AI (MANDATORY) ============
s = new_slide(); header(s, "Responsible AI Considerations", "Mandatory — Built Into the Design"); footer(s, 11)
pillars = [
    ("FAIRNESS", "Severity is normalised against a ward risk index, so wards with historically poorer sanitation are not deprioritised. Every score's factors are visible.", EMER),
    ("TRANSPARENCY", "The SERI engine is explainable — \"why this severity\" on every complaint. The AI assistant is grounded strictly in live database records; it never fabricates data.", TEAL),
    ("ETHICS", "AI only prioritises and informs. Humans verify, resolve and confirm at every decision point — no algorithmic penalties for citizens or officers.", AMBER),
    ("PRIVACY", "No personal data exposed publicly; complaint evidence is role-restricted; the audit log records actions, not personal profiles.", RGBColor(0xB4, 0x50, 0x36)),
]
positions = [(0.82, 1.85), (6.9, 1.85), (0.82, 4.35), (6.9, 4.35)]
for (tag, desc, accent), (x, y) in zip(pillars, positions):
    rect(s, x, y, 5.6, 2.3, WHITE)
    rect(s, x, y, 5.6, 0.13, accent)
    tf = tbox(s, x + 0.25, y + 0.28, 5.1, 1.9)
    para(tf, tag, size=15, bold=True, color=accent, first=True, after=6)
    para(tf, desc, size=12.5, color=INK, after=0)

# ============ SLIDE 12 — TOOLS & LEARNING JOURNEY ============
s = new_slide(); header(s, "Internship → Project: What I Applied", "IBM SkillsBuild Journey"); footer(s, 12)
rows = [
    ("IBM BOB", "Ideation & problem framing — refined the co-governance workflow before writing code.", EMER),
    ("Agentic AI Lab (Aug 2026)", "Tool usage, function calling, autonomous workflows — applied directly: our Python agent calls tools over the live complaints database.", TEAL),
    ("Guided Learning Path (GLE)", "AI fundamentals + responsible AI principles on IBM SkillsBuild — shaping the explainability and audit features.", GREEN),
    ("Mindset shift", "From \"calling models with prompts\" to building grounded, tool-using agents over real data.", AMBER),
]
y = 1.9
for tag, desc, accent in rows:
    rect(s, 0.82, y, 11.7, 1.05, WHITE)
    rect(s, 0.82, y, 0.14, 1.05, accent)
    tf = tbox(s, 1.15, y + 0.08, 3.4, 0.9, anchor=MSO_ANCHOR.MIDDLE)
    para(tf, tag, size=14, bold=True, color=accent, first=True, after=0)
    tf = tbox(s, 4.75, y + 0.08, 7.5, 0.9, anchor=MSO_ANCHOR.MIDDLE)
    para(tf, desc, size=12.5, color=INK, first=True, after=0)
    y += 1.22

# ============ SLIDE 13 — CONCLUSION ============
s = new_slide(GREEN_D)
rect(s, 0, 0, 13.333, 0.18, EMER)
tf = tbox(s, 0.9, 1.0, 11.5, 1.6)
para(tf, "Every complaint has a story, a status,", size=30, bold=True, color=WHITE, first=True, after=2)
para(tf, "and a closing signature from the citizen.", size=30, bold=True, color=EMER, after=0)
tf = tbox(s, 0.9, 3.0, 11.5, 2.4)
bullets(tf, [
    "Roadmap: citizen-officer comment threads · before/after photo galleries · multi-city rollout · GCC field pilot.",
    "Full source code published openly on GitHub.",
], size=15.5, gap=12, color=RGBColor(0xD9, 0xE8, 0xDF))
chip(s, 0.9, 5.35, 5.2, 0.6, "GitHub: github.com/lithes-coder/Sanitation-co-governance", fill=EMER, tcolor=GREEN_D, size=12.5)
tf = tbox(s, 0.9, 6.35, 11.5, 0.8)
para(tf, "Thank you — Questions?  ·  Lithesh S · Sri Sairam Institute of Technology", size=16, bold=True, color=WHITE, first=True, after=0)

OUT = os.path.join(HERE, "..", "SCGIP-Presentation.pptx")
prs.save(OUT)
print("Saved:", os.path.abspath(OUT), "slides:", len(prs.slides.__iter__.__self__._sldIdLst))
